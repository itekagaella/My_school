#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script pour générer le document de présentation complet du logiciel My_School au format .DOCX
"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import datetime

# Créer le document
doc = Document()

# Styles personnalisés
styles = doc.styles

# Titre
title = doc.add_heading('MY_SCHOOL - DOCUMENTATION DU LOGICIEL', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Sous-titre
subtitle = doc.add_paragraph()
subtitle_run = subtitle.add_run('Système de Gestion Scolaire')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run.font.size = Pt(16)
subtitle_run.font.italic = True

doc.add_paragraph()
date_para = doc.add_paragraph()
date_run = date_para.add_run(f'Généré le {datetime.now().strftime("%d/%m/%Y")}')
date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_run.font.size = Pt(10)

doc.add_page_break()

# Table des matières
doc.add_heading('Table des matières', 1)
toc_items = [
    '1. Présentation générale',
    '2. Objectifs du logiciel',
    '3. Prérequis techniques',
    '4. Installation',
    '5. Architecture technique',
    '6. Fonctionnalités principales',
    '7. Gestion des rôles et des utilisateurs',
    '8. Modules fonctionnels',
    '9. Sécurité',
    '10. Mise à jour et maintenance',
    '11. Sauvegarde et restauration',
    '12. Support et contact',
    '13. Annexes'
]
for item in toc_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# 1. Présentation générale
doc.add_heading('1. Présentation générale', 1)
doc.add_paragraph(
    'My_School est un système de gestion scolaire complet développé en PHP 8.5 avec une base de données '
    'PostgreSQL. Il permet de gérer l\'ensemble des activités administratives, pédagogiques et '
    'financières d\'un établissement scolaire.'
)
doc.add_paragraph(
    'Conçu pour être simple d\'utilisation, ce logiciel s\'adresse aux établissements scolaires de tous '
    'niveaux et offre une interface claire et intuitive pour chaque type d\'utilisateur.'
)

# 2. Objectifs du logiciel
doc.add_heading('2. Objectifs du logiciel', 1)
objectifs = [
    'Dématérialiser et centraliser la gestion scolaire',
    'Automatiser les tâches administratives répétitives',
    'Faciliter le suivi pédagogique des élèves',
    'Gérer efficacement les frais scolaires et les paiements',
    'Améliorer la communication entre l\'établissement, les élèves et leurs parents',
    'Assurer la traçabilité et la sécurité des données',
    'Générer des rapports et statistiques en temps réel'
]
for obj in objectifs:
    doc.add_paragraph(obj, style='List Bullet')

# 3. Prérequis techniques
doc.add_heading('3. Prérequis techniques', 1)
doc.add_heading('3.1. Environnement serveur', 2)
server_req = [
    'PHP 8.5 ou supérieur avec les extensions : pdo_pgsql, pgsql, gd, mbstring, openssl, curl, zip, fileinfo',
    'PostgreSQL 14 ou supérieur',
    'Serveur web Apache, Nginx ou serveur PHP intégré',
    'Espace disque minimum : 500 Mo',
    'Mémoire RAM minimum : 2 Go (recommandé 4 Go)'
]
for req in server_req:
    doc.add_paragraph(req, style='List Bullet')

doc.add_heading('3.2. Environnement client', 2)
client_req = [
    'Navigateur web moderne (Chrome, Firefox, Edge, Safari)',
    'Connexion Internet ou réseau local',
    'Résolution d\'écran minimum : 1024x768'
]
for req in client_req:
    doc.add_paragraph(req, style='List Bullet')

# 4. Installation
doc.add_heading('4. Installation', 1)
doc.add_heading('4.1. Installation avec Laragon (recommandée)', 2)
install_laragon = [
    'Installer Laragon (Version Full)',
    'Placer le dossier My_School dans C:\\laragon\\www\\',
    'Démarrer Apache et PostgreSQL depuis Laragon',
    'Ouvrir http://localhost/My_School/install.php',
    'Suivre l\'assistant d\'installation',
    'Supprimer install.php après installation réussie'
]
for step in install_laragon:
    doc.add_paragraph(step, style='List Number')

doc.add_heading('4.2. Installation avec XAMPP', 2)
install_xampp = [
    'Installer XAMPP et PostgreSQL',
    'Placer le dossier dans C:\\xampp\\htdocs\\My_School\\',
    'Activer les extensions pdo_pgsql, pgsql, gd, mbstring dans php.ini',
    'Démarrer Apache et PostgreSQL',
    'Accéder à http://localhost/My_School/install.php et suivre l\'assistant',
    'Supprimer install.php après installation'
]
for step in install_xampp:
    doc.add_paragraph(step, style='List Number')

doc.add_heading('4.3. Configuration manuelle', 2)
doc.add_paragraph(
    'Copier le fichier .env.example vers .env et configurer les paramètres de connexion '
    'à la base de données PostgreSQL ainsi que la clé de chiffrement.'
)

# 5. Architecture technique
doc.add_heading('5. Architecture technique', 1)
doc.add_heading('5.1. Technologies utilisées', 2)
tech_table = doc.add_table(rows=1, cols=2)
tech_table.style = 'Light Grid Accent 1'
hdr_cells = tech_table.rows[0].cells
hdr_cells[0].text = 'Composant'
hdr_cells[1].text = 'Technologie'

tech_data = [
    ['Langage backend', 'PHP 8.5'],
    ['Base de données', 'PostgreSQL 14+'],
    ['Frontend', 'HTML5, CSS3, JavaScript, Bootstrap'],
    ['Icônes', 'Font Awesome'],
    ['Graphiques', 'Chart.js'],
    ['Export PDF', 'TCPDF'],
    ['Export Excel', 'Générateur XLSX natif'],
]
for tech in tech_data:
    row_cells = tech_table.add_row().cells
    row_cells[0].text = tech[0]
    row_cells[1].text = tech[1]

doc.add_paragraph()

doc.add_heading('5.2. Structure des dossiers', 2)
structure = [
    'admin/ - Interface d\'administration',
    'client/ - Espace élèves/professeurs/personnel',
    'api/ - API REST',
    'config/ - Configuration, schéma SQL, migrations',
    'includes/ - Fonctions, authentification, mises en page',
    'assets/ - CSS, JavaScript, images, fichiers uploadés',
    'uploads/ - Documents et fichiers divers',
    'storage/ - Fichiers temporaires et mails de démo'
]
for s in structure:
    doc.add_paragraph(s, style='List Bullet')

# 6. Fonctionnalités principales
doc.add_heading('6. Fonctionnalités principales', 1)
doc.add_heading('6.1. Authentification et sécurité', 2)
auth_features = [
    'Connexion sécurisée avec sessions PHP',
    'Authentification à deux facteurs (2FA) par email',
    'Protection CSRF sur tous les formulaires',
    'Verrouillage automatique après tentatives échouées',
    'Déconnexion automatique en cas d\'inactivité',
    'Hachage bcrypt des mots de passe'
]
for f in auth_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('6.2. Gestion scolaire', 2)
scolaire_features = [
    'Gestion des élèves (inscription, classe, matricule)',
    'Gestion des professeurs et matières',
    'Emploi du temps avec créneaux standardisés',
    'Gestion des classes et sections',
    'Salles de classe et salles administratives'
]
for f in scolaire_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('6.3. Pédagogie', 2)
pedagogie_features = [
    'Saisie et consultation des notes',
    'Gestion des présences et absences',
    'Calcul automatique des moyennes',
    'Bulletins de notes',
    'Suivi de la discipline (incidents, sanctions)'
]
for f in pedagogie_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('6.4. Finances', 2)
finance_features = [
    'Gestion des frais scolaires',
    'Enregistrement des paiements (espèces uniquement)',
    'Génération de reçus',
    'Suivi des soldes par élève',
    'Rapports financiers'
]
for f in finance_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('6.5. Communication et documents', 2)
comm_features = [
    'Annonces et communications',
    'Chat interne entre utilisateurs',
    'Gestion documentaire avec versioning',
    'Partage de documents par classe/groupe',
    'Notifications en temps réel'
]
for f in comm_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('6.6. Rapports et statistiques', 2)
rapport_features = [
    'Rapports prédéfinis (élèves, professeurs, notes, présences)',
    'Rapports personnalisés avec sélection de colonnes',
    'Export PDF et Excel',
    'Tableaux de bord avec graphiques',
    'Statistiques globales'
]
for f in rapport_features:
    doc.add_paragraph(f, style='List Bullet')

# 7. Gestion des rôles et des utilisateurs
doc.add_heading('7. Gestion des rôles et des utilisateurs', 1)
doc.add_paragraph('Le système dispose de plusieurs rôles avec des permissions spécifiques :')

roles_table = doc.add_table(rows=1, cols=3)
roles_table.style = 'Light Grid Accent 1'
hdr = roles_table.rows[0].cells
hdr[0].text = 'Rôle'
hdr[1].text = 'Désignation'
hdr[2].text = 'Accès'

roles_data = [
    ['admin', 'Administrateur', 'Accès complet à toutes les fonctionnalités'],
    ['prof', 'Professeur', 'Gestion des notes, présences, emploi du temps'],
    ['eleve', 'Élève', 'Consultation (notes, emploi du temps, documents)'],
    ['surveillant', 'Surveillant', 'Suivi des élèves, présences, incidents'],
    ['finance', 'Gestionnaire financier', 'Frais, paiements, reçus (espèces uniquement)'],
    ['discipline', 'Responsable discipline', 'Incidents, sanctions, convocations'],
    ['personnel', 'Personnel (générique)', 'Accès limité (compatibilité)'],
]
for r in roles_data:
    row_cells = roles_table.add_row().cells
    row_cells[0].text = r[0]
    row_cells[1].text = r[1]
    row_cells[2].text = r[2]

doc.add_paragraph()

# 8. Modules fonctionnels
doc.add_heading('8. Modules fonctionnels', 1)
modules = [
    ('Espace Administrateur', 'Gestion complète du système, utilisateurs, paramètres, rapports, journal d\'activités'),
    ('Espace Professeur', 'Emploi du temps, saisie notes/présences, documents, communications'),
    ('Espace Élève', 'Consultation notes/bulletins, emploi du temps, paiements, documents, chat'),
    ('Espace Personnel', 'Modules selon le rôle (surveillance, finance, discipline)'),
]
for mod, desc in modules:
    p = doc.add_paragraph()
    p.add_run(mod + ' : ').bold = True
    p.add_run(desc)

# 9. Sécurité
doc.add_heading('9. Sécurité', 1)
security_measures = [
    'Authentification sécurisée avec sessions',
    'Authentification à deux facteurs (2FA)',
    'Protection CSRF sur tous les formulaires',
    'Requêtes préparées PDO (protection SQL injection)',
    'Échappement HTML (protection XSS)',
    'Hachage bcrypt des mots de passe',
    'Validation et nettoyage des entrées utilisateur',
    'Vérification des permissions par rôle',
    'Journalisation des activités',
    'Vérification du type réel des fichiers uploadés',
    'Clé de chiffrement pour les données sensibles'
]
for s in security_measures:
    doc.add_paragraph(s, style='List Bullet')

# 10. Mise à jour et maintenance
doc.add_heading('10. Mise à jour et maintenance', 1)
doc.add_paragraph(
    'Les mises à jour se font via le système de migrations automatique. Le fichier migrate.php '
    'permet d\'appliquer les nouvelles migrations sans affecter les données existantes.'
)
doc.add_heading('10.1. Procédure de mise à jour', 2)
update_steps = [
    'Sauvegarder la base de données (pg_dump)',
    'Sauvegarder les fichiers du projet',
    'Récupérer la nouvelle version',
    'Appliquer les migrations via migrate.php (navigateur ou CLI)',
    'Vérifier le bon fonctionnement'
]
for step in update_steps:
    doc.add_paragraph(step, style='List Number')

# 11. Sauvegarde et restauration
doc.add_heading('11. Sauvegarde et restauration', 1)
doc.add_heading('11.1. Sauvegarde de la base de données', 2)
doc.add_paragraph('Commande pg_dump :')
doc.add_paragraph('pg_dump -U postgres -d my_school > sauvegarde_$(date +%Y%m%d).sql', style='Intense Quote')

doc.add_heading('11.2. Restauration de la base de données', 2)
doc.add_paragraph('Commande psql :')
doc.add_paragraph('psql -U postgres -d my_school < sauvegarde.sql', style='Intense Quote')

doc.add_heading('11.3. Sauvegarde des fichiers', 2)
doc.add_paragraph(
    'Sauvegarder les dossiers assets/uploads/, uploads/, storage/ ainsi que le fichier .env '
    'qui contient les configurations sensibles.'
)

# 12. Support et contact
doc.add_heading('12. Support et contact', 1)
doc.add_paragraph(
    'Pour toute question, problème ou suggestion concernant l\'utilisation de My_School, '
    'veuillez contacter l\'administrateur système de votre établissement.'
)
doc.add_paragraph()
doc.add_paragraph('Version du logiciel : 1.0')
doc.add_paragraph('Dernière mise à jour : Octobre 2026')

# 13. Annexes
doc.add_heading('13. Annexes', 1)
doc.add_heading('13.1. Identifiants par défaut', 2)
default_table = doc.add_table(rows=1, cols=3)
default_table.style = 'Light Grid Accent 1'
dh = default_table.rows[0].cells
dh[0].text = 'Rôle'
dh[1].text = 'Email'
dh[2].text = 'Mot de passe'
default_row = default_table.add_row().cells
default_row[0].text = 'Administrateur'
default_row[1].text = 'admin@myschool.edu'
default_row[2].text = 'admin123'
doc.add_paragraph()
doc.add_paragraph(
    'Important : Changer ce mot de passe dès la première connexion.'
).runs[0].bold = True

doc.add_heading('13.2. Fichiers importants', 2)
files_list = [
    '.env - Configuration de la base de données et paramètres',
    'config/init.sql - Schéma complet de la base de données',
    'config/config.php - Configuration générale de l\'application',
    'includes/auth.php - Système d\'authentification',
    'migrate.php - Outil de mise à jour de la base de données',
    'install.php - Assistant d\'installation (à supprimer après usage)'
]
for f in files_list:
    doc.add_paragraph(f, style='List Bullet')

# Sauvegarder le document
output_path = '/home/iragame/My_School/DOCUMENTATION_MY_SCHOOL.docx'
doc.save(output_path)
print(f'Document créé avec succès : {output_path}')
print(f'Taille du fichier : {os.path.getsize(output_path)} octets')
