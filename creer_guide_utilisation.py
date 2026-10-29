#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Générateur du Guide d'Utilisation My_School
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os

doc = Document()

# Titre principal
title = doc.add_heading('GUIDE D\'UTILISATION', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
run = subtitle.add_run('MY_SCHOOL - Système de Gestion Scolaire')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run.font.size = Pt(18)
run.font.bold = True

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run(f'Version 1.0 | Généré le {datetime.now().strftime("%d/%m/%Y")}').font.size = Pt(10)

doc.add_page_break()

# Table des matières
doc.add_heading('Table des matières', 1)
toc = [
    '1. Introduction',
    '2. Accès à l\'application',
    '3. Interface générale',
    '4. Guide par rôle utilisateur',
    '5. Administration',
    '6. Gestion pédagogique',
    '7. Gestion financière',
    '8. Gestion de la discipline',
    '9. Communication',
    '10. Mon profil',
    '11. Notifications',
    '12. Sauvegarde et sécurité',
    '13. Dépannage et FAQ',
    '14. Annexe'
]
for item in toc:
    doc.add_paragraph(item, style='List Number')

doc.add_page_break()

# 1. Introduction
doc.add_heading('1. Introduction', 1)
doc.add_paragraph('My_School est un système complet de gestion scolaire permettant de gérer :')
for i in ['L\'administration et la gestion des utilisateurs', 'La structure scolaire (classes, matières, emplois du temps)', 'Le suivi pédagogique (notes, présences)', 'La gestion financière (frais, paiements)', 'La discipline scolaire (incidents, sanctions, convocations)', 'La communication (messages, documents, chat)', 'Les rapports et statistiques']:
    doc.add_paragraph(i, style='List Bullet')

# 2. Accès
doc.add_heading('2. Accès à l\'application', 1)
doc.add_heading('2.1. Connexion', 2)
doc.add_paragraph('Pour accéder à My_School :')
for i in ['Ouvrir votre navigateur web', 'Saisir l\'adresse : http://votre-serveur/My_School/', 'Entrer votre adresse email et mot de passe', 'Cliquer sur « Se connecter »']:
    doc.add_paragraph(i, style='List Number')
doc.add_heading('2.2. Première connexion', 2)
doc.add_paragraph('Email : admin@myschool.edu')
doc.add_paragraph('Mot de passe : admin123')
doc.add_paragraph('Il est fortement recommandé de changer ce mot de passe immédiatement depuis « Mon profil ».')

# 3. Interface
doc.add_heading('3. Interface générale', 1)
doc.add_heading('3.1. Barre latérale (Sidebar)', 2)
doc.add_paragraph('La barre latérale donne accès à tous les modules selon votre rôle. En haut se trouve votre profil (photo + nom + rôle).')
doc.add_heading('3.2. Barre supérieure (Topbar)', 2)
doc.add_paragraph('La barre supérieure affiche : le titre de la page, l\'icône notifications, l\'icône profil, l\'icône déconnexion.')
doc.add_heading('3.3. Tableau de bord', 2)
doc.add_paragraph('Chaque rôle dispose d\'un tableau de bord personnalisé avec des statistiques et raccourcis.')

# 4. Par rôle
doc.add_heading('4. Guide par rôle utilisateur', 1)
doc.add_heading('4.1. Administrateur', 2)
doc.add_paragraph('Accès complet à tous les modules depuis l\'espace Administration (/admin/).')
doc.add_heading('4.2. Élève', 2)
for i in ['Consulter ses notes', 'Consulter son emploi du temps', 'Consulter ses présences/absences', 'Consulter sa situation financière', 'Consulter discipline', 'Accéder aux clubs', 'Consulter documents', 'Utiliser le chat']:
    doc.add_paragraph(i, style='List Bullet')
doc.add_heading('4.3. Professeur', 2)
for i in ['Gérer ses élèves', 'Saisir et modifier les notes', 'Gérer les présences', 'Consulter emploi du temps', 'Envoyer des communications']:
    doc.add_paragraph(i, style='List Bullet')
doc.add_heading('4.4. Surveillant', 2)
for i in ['Suivi des élèves', 'Saisie des présences', 'Signalement d\'incidents disciplinaires']:
    doc.add_paragraph(i, style='List Bullet')
doc.add_heading('4.5. Gestionnaire Financier', 2)
for i in ['Gestion des frais scolaires', 'Enregistrement des paiements en espèces', 'Génération des reçus', 'Suivi des soldes', 'Rapports financiers']:
    doc.add_paragraph(i, style='List Bullet')
doc.add_heading('4.6. Responsable Discipline', 2)
for i in ['Gestion des incidents', 'Prononciation des sanctions', 'Gestion des convocations']:
    doc.add_paragraph(i, style='List Bullet')

# 5. Administration
doc.add_heading('5. Administration', 1)
doc.add_heading('5.1. Gestion des utilisateurs', 2)
doc.add_paragraph('Admin > Utilisateurs : créer, modifier, activer/désactiver les utilisateurs.')
doc.add_heading('5.2. Rôles et permissions', 2)
doc.add_paragraph('Admin > Rôles & Permissions : configurer les droits par rôle.')
doc.add_heading('5.3. Paramètres système', 2)
doc.add_paragraph('Admin > Paramètres : nom établissement, année scolaire (liste déroulante), coordonnées, session, uploads, seuil absences.')
doc.add_heading('5.4. Thèmes', 2)
doc.add_paragraph('Admin > Thèmes : choisir le thème (Par défaut, Sombre, Bleu, Vert).')
doc.add_heading('5.5. Sauvegarde', 2)
doc.add_paragraph('Admin > Sauvegarde : sauvegarde de la base PostgreSQL.')

# 6. Pédagogie
doc.add_heading('6. Gestion pédagogique', 1)
doc.add_heading('6.1. Classes', 2)
doc.add_paragraph('Création des classes avec niveau, section, année scolaire, capacité.')
doc.add_heading('6.2. Matières', 2)
doc.add_paragraph('Définition des matières avec code et coefficient.')
doc.add_heading('6.3. Emplois du temps', 2)
doc.add_paragraph('Horaires par classe, matière, professeur, jour, créneau. Publication possible.')
doc.add_heading('6.4. Notes', 2)
doc.add_paragraph('Saisie par type (interrogation, contrôle, devoir, composition) avec coefficient.')
doc.add_heading('6.5. Présences', 2)
doc.add_paragraph('Présent, Absent, Retard, Justifié par cours.')

# 7. Finances
doc.add_heading('7. Gestion financière', 1)
doc.add_heading('7.1. Frais scolaires', 2)
doc.add_paragraph('Définition par niveau, année scolaire, montant.')
doc.add_heading('7.2. Paiements', 2)
doc.add_paragraph('Paiements en espèces uniquement, reçu automatique avec numéro unique.')
doc.add_heading('7.3. Suivi des soldes', 2)
doc.add_paragraph('Total dû, payé, reste à payer par élève.')

# 8. Discipline
doc.add_heading('8. Gestion de la discipline', 1)
doc.add_heading('8.1. Incidents', 2)
doc.add_paragraph('Type, gravité (faible/moyenne/grave), lieu, description.')
doc.add_heading('8.2. Sanctions', 2)
doc.add_paragraph('Sanctions liées ou non à un incident avec dates.')
doc.add_heading('8.3. Convocations', 2)
doc.add_paragraph('Convocation parents/tuteurs avec date, lieu, motif.')

# 9. Communication
doc.add_heading('9. Communication', 1)
doc.add_heading('9.1. Communications', 2)
doc.add_paragraph('Messages à tous, par classe ou individuels avec priorité.')
doc.add_heading('9.2. Documents', 2)
doc.add_paragraph('Partage de documents avec gestion des téléchargements.')
doc.add_heading('9.3. Chat interne', 2)
doc.add_paragraph('Messagerie entre utilisateurs.')
doc.add_heading('9.4. Clubs scolaires', 2)
doc.add_paragraph('Gestion des clubs et inscriptions.')

# 10. Profil
doc.add_heading('10. Mon profil', 1)
doc.add_paragraph('Accessible depuis l\'icône profil en haut à droite.')
doc.add_paragraph('- Modifier informations personnelles')
doc.add_paragraph('- Changer/supprimer photo de profil (PNG, JPG, GIF, WebP)')
doc.add_paragraph('- Changer mot de passe (min 8 caractères)')
doc.add_paragraph('- Historique de connexions')

# 11. Notifications
doc.add_heading('11. Notifications', 1)
doc.add_paragraph('Badge rouge pour notifications non lues. Accessible via cloche.')

# 12. Sécurité
doc.add_heading('12. Sauvegarde et sécurité', 1)
for i in ['Sessions sécurisées avec expiration', 'Protection CSRF', 'Mots de passe bcrypt', '2FA optionnel', 'Verrouillage compte', 'Journal d\'activités', 'Contrôle par rôles/permissions', 'Sauvegardes']:
    doc.add_paragraph(i, style='List Bullet')

# 13. FAQ
doc.add_heading('13. Dépannage et FAQ', 1)
doc.add_heading('Mot de passe oublié ?', 2)
doc.add_paragraph('Utiliser « Mot de passe oublié » sur la page de connexion.')
doc.add_heading('Changer de thème ?', 2)
doc.add_paragraph('Admin > Thèmes uniquement.')
doc.add_heading('Modifier sa photo ?', 2)
doc.add_paragraph('Oui, depuis « Mon profil » pour tous les utilisateurs.')

# 14. Annexe
doc.add_heading('14. Annexe', 1)
doc.add_heading('14.1. Rôles', 2)
table = doc.add_table(rows=8, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text = 'Rôle'
table.rows[0].cells[1].text = 'Description'
roles = [('admin','Administrateur'),('eleve','Élève'),('prof','Professeur'),('surveillant','Surveillant'),('finance','Gestionnaire Financier'),('discipline','Responsable Discipline'),('personnel','Personnel')]
for i,(r,d) in enumerate(roles,1):
    table.rows[i].cells[0].text = r
    table.rows[i].cells[1].text = d

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('© 2026 My_School - Guide d\'utilisation').font.size = Pt(9)

filename = '/home/iragame/My_School/GUIDE_D_UTILISATION_MY_SCHOOL.docx'
doc.save(filename)
print(f'OK: {filename} ({os.path.getsize(filename)} bytes)')
