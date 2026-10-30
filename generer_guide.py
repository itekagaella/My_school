#!/usr/bin/env python3
import os, time, json, subprocess, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path('/home/iragame/My_School')
OUT = BASE / 'GUIDE_UTILISATION_MY_SCHOOL.docx'
CAP = BASE / 'guide_captures'
CAP.mkdir(exist_ok=True)

def take(url, name):
    out = CAP / f'{name}.png'
    cmd = [
        'chromium', '--headless=new', '--disable-gpu', '--no-sandbox', '--disable-dev-shm-usage',
        f'--window-size=1280,800', f'--screenshot={out}', url
    ]
    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    return out if out.exists() else None

doc = Document()
doc.add_heading('Guide d\'utilisation - My_School', 0)

doc.add_heading('1. Page de connexion', level=1)
doc.add_paragraph('Accès à l\'application via l\'écran de connexion. L\'utilisateur sélectionne son profil et saisit ses identifiants.')
img = take('http://localhost:8080/', '01_connexion')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('2. Tableau de bord Administrateur', level=1)
doc.add_paragraph('Après connexion en tant qu\'administrateur, le tableau de bord présente les statistiques globales de l\'établissement.')
img = take('http://localhost:8080/admin/index.php', '02_admin_dashboard')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('3. Gestion des élèves', level=1)
doc.add_paragraph('Module permettant d\'ajouter, modifier, consulter et archiver les élèves.')
img = take('http://localhost:8080/admin/eleves/index.php', '03_eleves')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('4. Gestion des professeurs', level=1)
doc.add_paragraph('Gestion du personnel enseignant avec leurs matières et classes.')
img = take('http://localhost:8080/admin/profs/index.php', '04_profs')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('5. Emplois du temps', level=1)
doc.add_paragraph('Création et consultation des emplois du temps par classe.')
img = take('http://localhost:8080/admin/horaires/index.php', '05_horaires')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('6. Notes', level=1)
doc.add_paragraph('Saisie et consultation des notes par les professeurs et administrateurs.')
img = take('http://localhost:8080/admin/notes/index.php', '06_notes')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('7. Présences', level=1)
doc.add_paragraph('Suivi des présences et absences des élèves.')
img = take('http://localhost:8080/admin/presences/index.php', '07_presences')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('8. Discipline', level=1)
doc.add_paragraph('Gestion des incidents, sanctions et convocations.')
img = take('http://localhost:8080/client/discipline/index.php', '08_discipline')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('9. Finance', level=1)
doc.add_paragraph('Gestion des frais scolaires, paiements et suivi des soldes.')
img = take('http://localhost:8080/admin/finance/index.php', '09_finance')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('10. Communications', level=1)
doc.add_paragraph('Envoi de communications aux élèves, parents et personnel.')
img = take('http://localhost:8080/admin/communications/index.php', '10_communications')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('11. Documents', level=1)
doc.add_paragraph('Gestion documentaire avec partage et versioning.')
img = take('http://localhost:8080/admin/documents/index.php', '11_documents')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('12. Archivage électronique (GED)', level=1)
doc.add_paragraph('Module d\'archivage des documents officiels depuis le module Documents avec référence, type, année scolaire et mots-clés.')
img = take('http://localhost:8080/admin/archivage-electronique/index.php', '12_archivage_electronique')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('13. Recherche documentaire', level=1)
doc.add_paragraph('Recherche unifiée dans les documents et l\'archivage électronique par mot-clé, titre, description ou référence.')
img = take('http://localhost:8080/admin/recherche-documentaire/index.php', '13_recherche_documentaire')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('14. Paramètres & Sauvegarde', level=1)
doc.add_paragraph('Configuration de l\'établissement, utilisateurs, rôles, permissions et sauvegarde de la base de données.')
img = take('http://localhost:8080/admin/parametres/index.php', '14_parametres')
if img: doc.add_picture(str(img), width=Inches(6))

doc.add_heading('15. Espace Élève', level=1)
doc.add_paragraph('Consultation des notes, emplois du temps, présences, communications et documents.')
img = take('http://localhost:8080/client/eleve/index.php', '15_espace_eleve')
if img: doc.add_picture(str(img), width=Inches(6))

doc.save(str(OUT))
print('done', OUT)
