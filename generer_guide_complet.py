#!/usr/bin/env python3
import os, time, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches

BASE = Path('/home/iragame/My_School')
OUT = BASE / 'GUIDE_UTILISATION_MY_SCHOOL.docx'
CAP = BASE / 'guide_captures'
CAP.mkdir(exist_ok=True)

def take(url, name, wait=3):
    out = CAP / f'{name}.png'
    cmd = [
        'chromium', '--headless=new', '--disable-gpu', '--no-sandbox', '--disable-dev-shm-usage',
        f'--window-size=1440,900', f'--screenshot={out}', url
    ]
    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    time.sleep(wait)
    return out if out.exists() else None

doc = Document()
doc.add_heading('Guide d\'utilisation - My_School', 0)
doc.add_paragraph('Ce guide présente les différentes interfaces de l\'application My_School avec leurs fonctionnalités.')

# Connexion
doc.add_heading('1. Authentification', level=1)
doc.add_paragraph('Page d\'accueil permettant de sélectionner le profil et de s\'authentifier.')
img = take('http://localhost:8080/', '01_auth')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('2. Tableau de bord Administrateur', level=1)
doc.add_paragraph('Vue d\'ensemble avec statistiques, effectifs, paiements, communications récentes.')
img = take('http://localhost:8080/admin/index.php', '02_admin_accueil')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('3. Gestion des élèves', level=1)
doc.add_paragraph('Liste des élèves, ajout, modification, consultation et archivage.')
img = take('http://localhost:8080/admin/eleves/index.php', '03_eleves_liste')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('4. Gestion des professeurs', level=1)
doc.add_paragraph('Gestion des enseignants, affectations matières/classes.')
img = take('http://localhost:8080/admin/profs/index.php', '04_profs_liste')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('5. Gestion du personnel', level=1)
doc.add_paragraph('Administration du personnel non-enseignant.')
img = take('http://localhost:8080/admin/personnel/index.php', '05_personnel')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('6. Classes et inscriptions', level=1)
doc.add_paragraph('Gestion des classes, niveaux et inscriptions des élèves.')
img = take('http://localhost:8080/admin/classes/index.php', '06_classes')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('7. Emplois du temps', level=1)
doc.add_paragraph('Création et publication des emplois du temps par classe.')
img = take('http://localhost:8080/admin/horaires/index.php', '07_horaires')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('8. Notes et évaluations', level=1)
doc.add_paragraph('Saisie des notes, bulletins et suivi des résultats.')
img = take('http://localhost:8080/admin/notes/index.php', '08_notes')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('9. Présences', level=1)
doc.add_paragraph('Appel, suivi et statistiques des présences/absences.')
img = take('http://localhost:8080/admin/presences/index.php', '09_presences')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('10. Discipline', level=1)
doc.add_paragraph('Incidents, sanctions, avertissements et convocations.')
img = take('http://localhost:8080/admin/discipline/index.php', '10_discipline')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('11. Finance - Frais scolaires', level=1)
doc.add_paragraph('Configuration des frais par classe et par niveau.')
img = take('http://localhost:8080/admin/finance/frais.php', '11_finance_frais')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('12. Finance - Paiements', level=1)
doc.add_paragraph('Enregistrement et suivi des paiements des élèves.')
img = take('http://localhost:8080/admin/finance/paiements.php', '12_finance_paiements')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('13. Communications', level=1)
doc.add_paragraph('Diffusion d\'informations à l\'ensemble de l\'établissement.')
img = take('http://localhost:8080/admin/communications/index.php', '13_communications')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('14. Clubs scolaires', level=1)
doc.add_paragraph('Gestion des activités extra-scolaires.')
img = take('http://localhost:8080/admin/clubs/index.php', '14_clubs')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('15. Gestion documentaire', level=1)
doc.add_paragraph('Bibliothèque de documents avec partage ciblé.')
img = take('http://localhost:8080/admin/documents/index.php', '15_documents')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('16. Archivage électronique (GED)', level=1)
doc.add_paragraph('Versement et classement des documents officiels. Seuls les documents existants dans le module Documents peuvent être archivés.')
img = take('http://localhost:8080/admin/archivage-electronique/index.php', '16_archivage_elec')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('17. Recherche documentaire', level=1)
doc.add_paragraph('Moteur de recherche unifié dans Documents et Archivage électronique.')
img = take('http://localhost:8080/admin/recherche-documentaire/index.php', '17_recherche_doc')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('18. Utilisateurs & Permissions', level=1)
doc.add_paragraph('Gestion des comptes utilisateurs, rôles et permissions granulaires.')
img = take('http://localhost:8080/admin/users/index.php', '18_utilisateurs')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('19. Paramètres de l\'établissement', level=1)
doc.add_paragraph('Configuration générale, logo, informations de l\'établissement.')
img = take('http://localhost:8080/admin/parametres/index.php', '19_parametres')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('20. Sauvegarde & Restauration', level=1)
doc.add_paragraph('Outil de sauvegarde automatique et manuelle de la base de données.')
img = take('http://localhost:8080/admin/parametres/backup.php', '20_sauvegarde')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('21. Journaux & Activité', level=1)
doc.add_paragraph('Traçabilité des actions effectuées dans l\'application.')
img = take('http://localhost:8080/admin/logs/index.php', '21_logs')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('22. Rapports & Statistiques', level=1)
doc.add_paragraph('Génération de rapports PDF/Excel et tableaux statistiques.')
img = take('http://localhost:8080/admin/rapports/index.php', '22_rapports')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('23. Espace Élève', level=1)
doc.add_paragraph('Tableau de bord personnel avec notes, emploi du temps, présences, communications.')
img = take('http://localhost:8080/client/eleve/index.php', '23_espace_eleve')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.add_heading('24. Espace Professeur', level=1)
doc.add_paragraph('Saisie des notes, gestion des présences, consultations.')
img = take('http://localhost:8080/client/prof/index.php', '24_espace_prof')
if img: doc.add_picture(str(img), width=Inches(6.5))

doc.save(str(OUT))
print('done')
