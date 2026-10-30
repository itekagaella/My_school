#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.shared import Inches

BASE = Path('/home/iragame/My_School')
OUT = BASE / 'GUIDE_UTILISATION_MY_SCHOOL.docx'
CAP = BASE / 'guide_captures'

doc = Document()
doc.add_heading('Guide d\'utilisation - My_School', 0)
doc.add_paragraph('Version complète avec captures réelles de toutes les interfaces.')
doc.add_paragraph('Date : Octobre 2025')

sections = [
    ('1. Page d\'accueil / Authentification', ['01_accueil.png'], 'Accès à l\'application My_School avec sélection du profil de connexion.'),
    ('2. Tableau de bord Administrateur', ['02_admin_accueil.png'], 'Vue synthétique : effectifs, inscriptions, paiements, communications.'),
    ('3. Gestion des élèves', ['03_eleves.png'], 'Liste, ajout, modification, consultation et archivage des élèves.'),
    ('4. Gestion des professeurs', ['04_profs.png'], 'Enseignants, matières affectées, emplois du temps.'),
    ('5. Gestion du personnel', ['05_personnel.png'], 'Personnel administratif et technique.'),
    ('6. Classes', ['06_classes.png'], 'Création des classes, répartition des élèves.'),
    ('7. Emplois du temps', ['07_horaires.png'], 'Planification hebdomadaire par classe.'),
    ('8. Notes', ['08_notes.png'], 'Saisie des notes par évaluation, consultation des bulletins.'),
    ('9. Présences', ['09_presences.png'], 'Appel quotidien et suivi des absences.'),
    ('10. Discipline', ['10_discipline.png'], 'Incidents, sanctions, avertissements, convocations.'),
    ('11. Finance - Tableau de bord', ['11_finance.png'], 'Situation financière globale.'),
    ('12. Finance - Frais scolaires', ['12_finance_frais.png'], 'Paramétrage des frais par classe.'),
    ('13. Finance - Paiements', ['13_finance_paiements.png'], 'Enregistrement, reçus, suivi des soldes.'),
    ('14. Communications', ['14_communications.png'], 'Diffusion de messages à l\'établissement.'),
    ('15. Clubs scolaires', ['15_clubs.png'], 'Gestion des clubs et inscriptions.'),
    ('16. Documents', ['16_documents.png'], 'Bibliothèque documentaire partagée.'),
    ('17. Archivage électronique (GED)', ['17_archivage.png'], 'Archivage des documents officiels depuis le module Documents.'),
    ('18. Recherche documentaire', ['18_recherche.png'], 'Recherche unifiée dans Documents et Archivage électronique.'),
    ('19. Utilisateurs et permissions', ['19_utilisateurs.png'], 'Comptes, rôles et droits d\'accès.'),
    ('20. Paramètres de l\'établissement', ['20_parametres.png'], 'Configuration générale, logo, informations.'),
    ('21. Sauvegarde', ['21_sauvegarde.png'], 'Sauvegarde et restauration de la base de données.'),
    ('22. Journaux d\'activité', ['22_logs.png'], 'Traçabilité des actions.'),
    ('23. Rapports et statistiques', ['23_rapports.png'], 'Rapports PDF/Excel, graphiques.'),
    ('24. Espace Élève - Accueil', ['24_eleve_accueil.png'], 'Tableau de bord élève.'),
    ('25. Espace Élève - Notes', ['25_eleve_notes.png'], 'Consultation des notes et bulletins.'),
    ('26. Espace Élève - Emplois du temps', ['26_eleve_horaires.png'], 'Emploi du temps personnel.'),
    ('27. Espace Élève - Présences', ['27_eleve_presences.png'], 'Suivi des présences et absences.'),
    ('28. Espace Élève - Discipline', ['28_eleve_disciplines.png'], 'Historique disciplinaire.'),
    ('29. Espace Élève - Communications', ['29_eleve_communications.png'], 'Messages reçus.'),
    ('30. Espace Élève - Documents', ['30_eleve_documents.png'], 'Documents partagés.'),
    ('31. Espace Élève - Finance', ['31_eleve_finance.png'], 'Frais et paiements.'),
    ('32. Espace Élève - Clubs', ['32_eleve_clubs.png'], 'Inscription aux clubs.'),
    ('33. Espace Professeur - Accueil', ['33_prof_accueil.png'], 'Tableau de bord professeur.'),
    ('34. Espace Professeur - Saisie des notes', ['34_prof_notes.png'], 'Saisie et modification des notes.'),
    ('35. Espace Professeur - Appel', ['35_prof_presences.png'], 'Prise des présences.'),
    ('36. Espace Professeur - Emplois du temps', ['36_prof_horaires.png'], 'Horaires du professeur.'),
    ('37. Espace Professeur - Communications', ['37_prof_communications.png'], 'Communications.'),
    ('38. Espace Professeur - Documents', ['38_prof_documents.png'], 'Documents partagés.'),
]

for title, imgs, desc in sections:
    doc.add_heading(title, level=1)
    doc.add_paragraph(desc)
    for img in imgs:
        p = CAP / img
        if p.exists():
            doc.add_picture(str(p), width=Inches(6.5))
        else:
            doc.add_paragraph(f'[Capture manquante: {img}]')

doc.save(str(OUT))
print('done')
